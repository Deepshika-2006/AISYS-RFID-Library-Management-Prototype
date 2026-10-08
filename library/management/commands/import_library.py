from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook
from library.models import Book,MigrationRun
class Command(BaseCommand):
 help='Controlled spreadsheet import with validation, duplicate detection and rollback.'
 def add_arguments(self,p): p.add_argument('path')
 def handle(self,*a,**o):
  path=o['path']; run=MigrationRun.objects.create(source_name=path)
  errors=[]; imported=[]
  try:
   wb=load_workbook(path,read_only=True,data_only=True); ws=wb.active; rows=list(ws.iter_rows(values_only=True)); run.total_rows=max(0,len(rows)-1)
   headers=[str(x).strip().lower() if x is not None else '' for x in rows[0]] if rows else []
   required=['accession_no','title','isbn','author','category','location','rfid_tag']
   missing=[h for h in required if h not in headers]
   if missing: raise ValueError('Missing columns: '+', '.join(missing))
   idx={h:headers.index(h) for h in required}; seen=set()
   with transaction.atomic():
    for n,row in enumerate(rows[1:],start=2):
     acc=str(row[idx['accession_no']] or '').strip(); title=str(row[idx['title']] or '').strip()
     if not acc or not title: run.rejected_rows+=1; errors.append({'row':n,'error':'accession_no and title are required'}); continue
     if acc in seen or Book.objects.filter(accession_no=acc).exists(): run.duplicate_rows+=1; errors.append({'row':n,'error':'duplicate accession_no'}); continue
     seen.add(acc); run.valid_rows+=1
     imported.append(Book(accession_no=acc,title=title,isbn=str(row[idx['isbn']] or ''),author=str(row[idx['author']] or ''),category=str(row[idx['category']] or ''),location=str(row[idx['location']] or 'Main Shelf'),rfid_tag=str(row[idx['rfid_tag']] or '') or None))
    Book.objects.bulk_create(imported); run.imported_rows=len(imported)
  except Exception as e:
   run.error_log=errors+[{'fatal':str(e)}]; run.save(); raise
  run.error_log=errors; run.save(); self.stdout.write(self.style.SUCCESS(f'Imported {run.imported_rows}; rejected {run.rejected_rows}; duplicates {run.duplicate_rows}'))
